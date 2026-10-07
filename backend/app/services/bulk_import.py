"""Versioned XLSX export and atomic, reviewed imports."""
import hashlib
import hmac
import json
import math
from io import BytesIO
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlmodel import Session, select

from app.config import settings
from app.models import Product, Service, ServiceCategory
from app.routers.catalog import ServiceInput
from app.schemas import ProductCreate

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
COLUMNS = {
    "products": [("id", "ID"), ("sku", "SKU"), ("name", "Nombre"), ("category", "Categoría"),
                 ("unit", "Unidad"), ("stock_quantity", "Stock"), ("minimum_quantity", "Stock mínimo"),
                 ("cost_clp", "Costo CLP"), ("price_clp", "Precio CLP"), ("active", "Activo"), ("version", "_version")],
    "services": [("id", "ID"), ("code", "Código"), ("name", "Nombre"), ("category", "Categoría"),
                 ("description", "Descripción"), ("version", "_version")],
}


def model_for(kind):
    if kind not in COLUMNS:
        raise HTTPException(404, "Tipo de planilla no permitido")
    return Product if kind == "products" else Service


def snapshot(row, categories):
    value = row.model_dump()
    if isinstance(row, Service):
        value["category"] = categories[row.category_id].name
    return value


def version(kind, value):
    data = json.dumps([kind, value], sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hmac.new(settings.jwt_secret.encode(), data.encode(), hashlib.sha256).hexdigest()


def export_workbook(session: Session, kind: str):
    model = model_for(kind)
    rows = session.exec(select(model).order_by(model.id).limit(MAX_ROWS + 1)).all()
    if len(rows) > MAX_ROWS:
        raise HTTPException(422, f"El catálogo supera el máximo de {MAX_ROWS} registros por planilla")
    categories = {c.id: c for c in session.exec(select(ServiceCategory)).all()} if kind == "services" else {}
    book = Workbook()
    sheet = book.active
    sheet.title = "Datos"
    sheet.append([label for _, label in COLUMNS[kind]])
    for row in rows:
        data = snapshot(row, categories)
        data["version"] = version(kind, data)
        if kind == "products":
            data["active"] = "Sí" if row.active else "No"
        sheet.append([data.get(key) for key, _ in COLUMNS[kind]])
        # Export names/codes as literal text, never executable spreadsheet formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="C00000")
    for i, (key, _) in enumerate(COLUMNS[kind], 1):
        sheet.column_dimensions[get_column_letter(i)].width = 42 if key in {"name", "description"} else 20
    sheet.column_dimensions[get_column_letter(len(COLUMNS[kind]))].hidden = True
    if kind == "products":
        validation = DataValidation(type="list", formula1='"Sí,No"')
        sheet.add_data_validation(validation)
        validation.add(f"J2:J{MAX_ROWS + 1}")
    instructions = book.create_sheet("Instrucciones")
    for text in [
        "Carga masiva Gudex — " + kind,
        "Edita la hoja Datos. Conserva encabezados, ID y la columna oculta _version de los registros existentes.",
        "Para agregar registros, deja ID y _version vacíos. No reutilices códigos/SKU existentes.",
        "Quitar filas no elimina registros. En productos, Activo=No archiva el producto.",
        "Nombre obligatorio. Servicios: Código y Categoría obligatorios. Productos: Unidad obligatoria.",
        "Stock y Stock mínimo: cantidades numéricas >= 0. Costo y Precio CLP: números enteros >= 0, sin símbolos.",
        "Los cambios de Stock generan ajustes con referencia a la importación.",
        "Las categorías nuevas de servicios se crean solo al confirmar la vista previa.",
        "Si un registro cambió desde la descarga, descarga otra planilla para resolver el conflicto.",
        f"Máximo {MAX_ROWS} filas y 5 MB. No se admiten fórmulas. Guarda como Excel .xlsx.",
        "Sube la planilla, revisa todos los cambios y confirma. Si hay errores, no se guarda ninguna fila.",
    ]:
        instructions.append([text])
    instructions.column_dimensions["A"].width = 125
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def read_workbook(content: bytes, kind: str):
    model_for(kind)
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(413, "La planilla está vacía o supera 5 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise HTTPException(413, "La planilla descomprimida supera el límite permitido")
            if any('vbaproject' in i.filename.lower() for i in archive.infolist()):
                raise HTTPException(422, "No se admiten macros")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if "Datos" not in book.sheetnames:
                raise ValueError("Falta la hoja Datos")
            sheet = book["Datos"]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            headers = next(iterator, ())
            if [c.value for c in headers] != [label for _, label in COLUMNS[kind]]:
                raise ValueError("Los encabezados no coinciden. Usa la planilla descargada de este módulo")
            result = []
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1:
                    raise ValueError(f"La planilla supera {MAX_ROWS} filas")
                if all(c.value is None for c in cells):
                    continue
                if len(cells) > len(headers) and any(c.value is not None for c in cells[len(headers):]):
                    raise ValueError(f"Fila {number}: hay columnas adicionales")
                values = {key: cells[i].value if i < len(cells) else None for i, (key, _) in enumerate(COLUMNS[kind])}
                formulas = [COLUMNS[kind][i][1] for i, c in enumerate(cells[:len(headers)]) if c.data_type in {"f", "e"}]
                result.append({"row": number, "values": values, "formulas": formulas})
            if not result:
                raise ValueError("La planilla no contiene registros")
            # Reject dates and other unsupported cell types before persisting JSON.
            for row in result:
                if any(v is not None and not isinstance(v, (str, int, float, bool)) for v in row["values"].values()):
                    raise ValueError(f"Fila {row['row']}: hay fechas o valores no compatibles")
            return result
        finally:
            book.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, str(exc) if isinstance(exc, ValueError) else "Archivo Excel inválido. Usa una planilla .xlsx")


def text(value, label, maximum, required=False):
    if value is None:
        value = ""
    if not isinstance(value, str):
        # Preserve numeric-looking SKUs only when explicitly entered as text.
        raise ValueError(f"{label}: usa texto (códigos con ceros iniciales deben tener formato Texto)")
    value = value.strip()
    if required and not value:
        raise ValueError(f"{label}: campo obligatorio")
    if len(value) > maximum:
        raise ValueError(f"{label}: máximo {maximum} caracteres")
    return value


def number(value, label, integer=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError(f"{label}: ingresa un número no negativo")
    if integer and (value != int(value) or value > 2147483647):
        raise ValueError(f"{label}: ingresa un entero entre 0 y 2147483647")
    return int(value) if integer else float(value)


def normalize(raw, kind):
    data = {"name": text(raw["name"], "Nombre", 120 if kind == "products" else 160, True),
            "category": text(raw["category"], "Categoría", 80, kind == "services")}
    if kind == "services":
        data.update(code=text(raw["code"], "Código", 40, True).upper(),
                    description=text(raw["description"], "Descripción", 2000))
        ServiceInput.model_validate({**{k: v for k, v in data.items() if k != "category"}, "category_id": 1})
    else:
        active = raw["active"]
        if isinstance(active, bool):
            active_value = active
        elif isinstance(active, str) and active.strip().casefold() in {"sí", "si", "no"}:
            active_value = active.strip().casefold() != "no"
        else:
            raise ValueError("Activo: indica Sí o No")
        data.update(sku=text(raw["sku"], "SKU", 40).upper() or None,
                    category=data["category"] or None, unit=text(raw["unit"], "Unidad", 24, True),
                    stock_quantity=number(raw["stock_quantity"], "Stock"),
                    minimum_quantity=number(raw["minimum_quantity"], "Stock mínimo"),
                    cost_clp=number(raw["cost_clp"], "Costo CLP", True),
                    price_clp=number(raw["price_clp"], "Precio CLP", True), active=active_value)
        ProductCreate.model_validate(data)
    return data


def plan_import(session, kind, rows, lock=False):
    model = model_for(kind)
    statement = select(model).order_by(model.id)
    category_statement = select(ServiceCategory).order_by(ServiceCategory.id)
    if lock:
        statement = statement.with_for_update()
        category_statement = category_statement.with_for_update()
    current = {r.id: r for r in session.exec(statement).all()}
    categories = {c.id: c for c in session.exec(category_statement).all()} if kind == "services" else {}
    category_names = {c.name.casefold(): c.name for c in categories.values()}
    code_field = "sku" if kind == "products" else "code"
    codes = {}
    for item in current.values():
        code = getattr(item, code_field)
        if code:
            codes.setdefault(code.strip().upper(), set()).add(item.id)
    seen_ids, seen_codes, new_categories = set(), set(), set()
    errors, changes = [], []
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    for source in rows:
        raw, row_number = source["values"], source["row"]
        try:
            if source["formulas"]:
                raise ValueError("No se admiten fórmulas o errores en: " + ", ".join(source["formulas"]))
            row_id = None if raw["id"] in (None, "") else number(raw["id"], "ID", True)
            if row_id is not None:
                if row_id in seen_ids:
                    raise ValueError("ID repetido en la planilla")
                seen_ids.add(row_id)
                if row_id not in current:
                    raise ValueError("ID inexistente; para crear un registro deja ID y _version vacíos")
                original = snapshot(current[row_id], categories)
                if (not isinstance(raw["version"], str) or len(raw["version"]) != 64
                        or any(c not in "0123456789abcdef" for c in raw["version"])
                        or not hmac.compare_digest(raw["version"], version(kind, original))):
                    raise ValueError("Conflicto: registro modificado desde la descarga o referencia _version alterada; descarga una nueva planilla")
            else:
                if raw["version"] not in (None, ""):
                    raise ValueError("Fila nueva: deja ID y _version vacíos")
                original = None
            data = normalize(raw, kind)
            code = data[code_field]
            if code:
                if code in seen_codes:
                    raise ValueError("Código/SKU repetido en la planilla")
                seen_codes.add(code)
                if codes.get(code, set()) - {row_id}:
                    raise ValueError("Código/SKU ya asignado a otro registro; conserva su ID para actualizarlo")
            if kind == "services":
                key = data["category"].casefold()
                if key not in category_names:
                    new_categories.add(data["category"])
                    category_names[key] = data["category"]
                data["category"] = category_names[key]
            fields = list(data) if original is None else [k for k, v in data.items() if original.get(k) != v]
            action = "created" if original is None else "updated" if fields else "unchanged"
            counts[action] += 1
            changes.append({"row": row_number, "id": row_id, "action": action, "name": data["name"],
                            "fields": fields, "before": {k: original.get(k) for k in fields} if original else {},
                            "data": data})
        except (ValueError, ValidationError) as exc:
            errors.append({"row": row_number, "message": str(exc)})
    return {**counts, "new_categories": sorted(new_categories), "errors": errors, "changes": changes}
