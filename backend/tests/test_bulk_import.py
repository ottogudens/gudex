from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select

from app.database import engine
from app.main import app
from app.models import BulkImport, Product, Service, ServiceCategory, StockMovement, User, UserRole
from app.security import create_access_token


@pytest.fixture
def client():
    with TestClient(app) as client:
        with Session(engine) as session:
            user = User(email=f'bulk-{uuid4()}@test.cl', full_name='Import admin', password_hash='unused', role=UserRole.admin)
            session.add(user); session.commit(); session.refresh(user)
            client.headers['Authorization'] = 'Bearer ' + create_access_token(user)
        yield client


def workbook(client, kind='products'):
    response = client.get(f'/api/v1/bulk/{kind}/export')
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


def upload(client, book, kind='products'):
    buf = BytesIO(); book.save(buf)
    response = client.post(f'/api/v1/bulk/{kind}/preview', files={'file': ('datos.xlsx', buf.getvalue())})
    assert response.status_code == 200, response.text
    return response.json()


def product(client, **kwargs):
    data = {'name': 'Producto planilla', 'sku': uuid4().hex, 'stock_quantity': 10, **kwargs}
    response = client.post('/api/v1/products', json=data)
    assert response.status_code == 201, response.text
    return response.json()


def only_row(book, row_id):
    sheet = book['Datos']
    values = next(list(row) for row in sheet.values if row[0] == row_id)
    sheet.delete_rows(2, sheet.max_row)
    sheet.append(values)
    return sheet


def confirm(client, preview, kind='products'):
    assert not preview['errors'], preview
    return client.post(f'/api/v1/bulk/{kind}/{preview["import_id"]}/confirm')


def test_export_update_new_stock_audit_and_idempotency(client):
    original = product(client)
    untouched = product(client, name='No incluido')
    book = workbook(client); sheet = only_row(book, original['id'])
    sheet.cell(2, 3, 'Nombre actualizado'); sheet.cell(2, 6, 14.5)
    sku = uuid4().hex
    sheet.append([None, sku, 'Producto nuevo', 'Filtros', 'unidad', 3, 1, 200, 500, 'Sí', None])
    preview = upload(client, book)
    assert (preview['created'], preview['updated']) == (1, 1)
    with Session(engine) as s:
        assert s.get(Product, original['id']).stock_quantity == 10
        assert s.exec(select(Product).where(Product.sku == sku)).first() is None
    result = confirm(client, preview)
    assert result.status_code == 200, result.text
    assert confirm(client, preview).json() == result.json()
    with Session(engine) as s:
        assert s.get(Product, original['id']).stock_quantity == 14.5
        assert s.get(Product, untouched['id']).active
        movements = s.exec(select(StockMovement).where(StockMovement.reference == preview['import_id'])).all()
        assert sorted(m.quantity_change for m in movements) == [3, 4.5]
        audit = s.get(BulkImport, preview['import_id'])
        assert audit.status == 'imported' and audit.confirmed_at
    assert client.get('/api/v1/bulk/products/history').json()[0]['id'] == preview['import_id']


def test_invalid_rows_prevent_entire_import(client):
    p = product(client)
    book = workbook(client); sheet = only_row(book, p['id'])
    sheet.cell(2, 3, 'No guardar')
    sheet.append([None, uuid4().hex, 'Precio malo', '', 'unidad', 1, 0, 2.5, -1, 'Sí', None])
    preview = upload(client, book)
    assert preview['import_id'] is None and preview['errors'][0]['row'] == 3
    with Session(engine) as s:
        assert s.get(Product, p['id']).name == p['name']


def test_stale_export_and_change_after_preview(client):
    p = product(client)
    book = workbook(client); sheet = only_row(book, p['id']); sheet.cell(2, 6, 20)
    preview = upload(client, book)
    client.post(f'/api/v1/products/{p["id"]}/stock-movements', json={'quantity_change': -1, 'reason': 'sale'})
    assert confirm(client, preview).status_code == 409
    stale = upload(client, book)
    assert stale['import_id'] is None and 'Conflicto' in stale['errors'][0]['message']
    with Session(engine) as s:
        assert s.get(Product, p['id']).stock_quantity == 9


def test_services_new_category_update_and_code_duplicate(client):
    cat = client.post('/api/v1/service-categories', json={'name': f'Cat {uuid4()}'}).json()
    service = client.post('/api/v1/services', json={'code': uuid4().hex, 'name': 'Anterior', 'category_id': cat['id']}).json()
    book = workbook(client, 'services'); sheet = only_row(book, service['id'])
    new_category = 'Nueva ' + uuid4().hex
    sheet.cell(2, 3, 'Editado'); sheet.cell(2, 4, new_category)
    sheet.append([None, uuid4().hex, 'Nuevo', new_category.lower(), 'Detalle', None])
    preview = upload(client, book, 'services')
    assert preview['new_categories'] == [new_category]
    assert confirm(client, preview, 'services').status_code == 200
    with Session(engine) as s:
        item = s.get(Service, service['id'])
        assert item.name == 'Editado'
        assert s.get(ServiceCategory, item.category_id).name == new_category
    book = workbook(client, 'services'); sheet = only_row(book, service['id'])
    sheet.append([None, service['code'], 'Duplicado', new_category, '', None])
    assert upload(client, book, 'services')['errors']


def test_archive_and_export_include_inactive_formula_as_literal(client):
    p = product(client, name='=SUM(1,2)', sku='000123-' + uuid4().hex[:8])
    book = workbook(client); sheet = only_row(book, p['id'])
    assert sheet.cell(2, 3).data_type == 'f'  # openpyxl assignment in test reinterprets text
    sheet.cell(2, 3).data_type = 's'
    sheet.cell(2, 10, 'No')
    assert confirm(client, upload(client, book)).status_code == 200
    exported = workbook(client)
    row = next(row for row in exported['Datos'] if row[0].value == p['id'])
    assert row[2].data_type == 's' and row[2].value == '=SUM(1,2)'
    assert row[9].value == 'No' and row[1].value == p['sku']


def test_duplicate_ids_formula_unknown_id_and_wrong_headers(client):
    p = product(client)
    book = workbook(client); sheet = only_row(book, p['id'])
    sheet.append([c.value for c in sheet[2]])
    assert 'ID repetido' in upload(client, book)['errors'][0]['message']
    sheet.delete_rows(3); sheet.cell(2, 8, '=1+2')
    assert 'fórmulas' in upload(client, book)['errors'][0]['message']
    sheet.cell(2, 8, 0); sheet.cell(2, 11, 'ñ' * 64)
    assert 'Conflicto' in upload(client, book)['errors'][0]['message']
    sheet.cell(2, 1, 999999)
    assert 'ID inexistente' in upload(client, book)['errors'][0]['message']
    sheet.cell(1, 1, 'Otro')
    buf = BytesIO(); book.save(buf)
    assert client.post('/api/v1/bulk/products/preview', files={'file': ('datos.xlsx', buf.getvalue())}).status_code == 422
    assert client.post('/api/v1/bulk/products/preview', files={'file': ('datos.xlsx', b'bad')}).status_code == 422


def test_expiry_owner_and_permissions(client):
    p = product(client)
    book = workbook(client); only_row(book, p['id']); preview = upload(client, book)
    with Session(engine) as s:
        record = s.get(BulkImport, preview['import_id']); record.expires_at = datetime.now(timezone.utc) - timedelta(minutes=1)
        s.add(record); s.commit()
    assert confirm(client, preview).status_code == 409
    with Session(engine) as s:
        other = User(email=f'other-{uuid4()}@test.cl', full_name='Other', password_hash='unused', role=UserRole.admin)
        s.add(other); s.commit(); s.refresh(other)
        client.headers['Authorization'] = 'Bearer ' + create_access_token(other)
    assert confirm(client, preview).status_code == 404
    with Session(engine) as s:
        other.role = UserRole.mechanic; s.add(other); s.commit(); s.refresh(other)
        client.headers['Authorization'] = 'Bearer ' + create_access_token(other)
    assert client.get('/api/v1/bulk/products/export').status_code == 403
    assert client.post('/api/v1/bulk/services/preview', files={'file': ('test.xlsx', b'bad')}).status_code == 403
    assert confirm(client, preview).status_code == 403
    assert client.get('/api/v1/bulk/products/history').status_code == 403


def test_new_code_added_after_preview_rolls_back(client):
    p = product(client)
    book = workbook(client); sheet = only_row(book, p['id']); sheet.cell(2, 3, 'No guardar')
    sku = uuid4().hex
    sheet.append([None, sku, 'Nuevo', '', 'unidad', 0, 0, 0, 0, 'Sí', None])
    preview = upload(client, book)
    product(client, sku=sku)
    assert confirm(client, preview).status_code == 409
    with Session(engine) as s:
        assert s.get(Product, p['id']).name == p['name']


def test_database_failure_rolls_back_updates_stock_and_audit(client):
    from sqlalchemy import event
    from sqlalchemy.exc import IntegrityError
    p = product(client)
    book = workbook(client); sheet = only_row(book, p['id']); sheet.cell(2, 6, 50)
    sheet.append([None, uuid4().hex, 'Trigger failure', '', 'unidad', 7, 0, 0, 0, 'Sí', None])
    preview = upload(client, book)
    def fail_insert(mapper, connection, target):
        if target.name == 'Trigger failure':
            raise IntegrityError('insert', {}, Exception('simulated constraint violation'))
    event.listen(Product, 'before_insert', fail_insert)
    try:
        assert confirm(client, preview).status_code == 409
    finally:
        event.remove(Product, 'before_insert', fail_insert)
    with Session(engine) as s:
        assert s.get(Product, p['id']).stock_quantity == 10
        assert s.exec(select(StockMovement).where(StockMovement.reference == preview['import_id'])).first() is None
        assert s.get(BulkImport, preview['import_id']).status == 'preview'
